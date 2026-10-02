# NOTE (release copy): absolute local paths were replaced with the placeholders
# <PROJECT_ROOT>, <DATA_ROOT>, <TOOLS_ROOT>, <HOME>, <PYTHON> etc. Adapt these at the
# top of the script (or via environment variables) before running. See README.md.
# -*- coding: utf-8 -*-
"""
build_tables2.py — 汇总 Supplementary Table S3（xlsx + tsv）与规范命名附表（最终数据版本）
输出：
  06_tables/Supplementary_Table_S3_junction_read_support.xlsx / .tsv
  06_tables/junction_reference_manifest.tsv
  06_tables/mapping_model_comparison.tsv
  06_tables/ambiguous_and_negative_reads.tsv
  06_tables/chromosome_state_and_copy_number.tsv
"""
import os
from collections import defaultdict

ATTJ = r'<PROJECT_ROOT>\longread_validation'
SR = os.path.join(ATTJ, '03_supporting_reads')
CT = os.path.join(ATTJ, '04_controls')
T = os.path.join(ATTJ, '06_tables')
REFDIR = os.path.join(ATTJ, '01_references')


def load_tsv(p):
    rows = []
    if not os.path.exists(p):
        return rows
    lines = [l for l in open(p, encoding='utf-8').read().splitlines() if l.strip()]
    hdr = None
    for l in lines:
        if l.startswith('#') and '\t' not in l:
            continue          # 纯注释行（说明文字）跳过
        if hdr is None:
            hdr = l.lstrip('#').split('\t')
            continue
        if l.startswith('#'):
            continue
        rows.append(dict(zip(hdr, l.split('\t'))))
    return rows


summ = load_tsv(os.path.join(SR, 'junction_support_summary.tsv'))
cand = load_tsv(os.path.join(SR, 'junction_candidates2.tsv'))
cls3 = load_tsv(os.path.join(SR, 'read_classification3.tsv'))
sjc = load_tsv(os.path.join(CT, 'state_junction_counts.tsv'))
cn = load_tsv(os.path.join(CT, 'copy_number_corrected.tsv'))
dbs = load_tsv(os.path.join(CT, 'depth_ratio_bootstrap3.tsv'))
drs = load_tsv(os.path.join(CT, 'depth_region_stats3.tsv'))
dsam = load_tsv(os.path.join(CT, 'depth_samtools_regions.tsv'))
mm = load_tsv(os.path.join(CT, 'multimapping_stats.tsv'))
mix = load_tsv(os.path.join(CT, 'mixed_structure_counts.tsv'))
r1 = load_tsv(os.path.join(CT, 'unit_1bp_withinread2.tsv'))
r1p = load_tsv(os.path.join(CT, 'unit_1bp_withinread_pairs2.tsv'))
r3 = load_tsv(os.path.join(CT, 'unit_1bp_withinread3.tsv'))
r3p = load_tsv(os.path.join(CT, 'unit_1bp_withinread_pairs3.tsv'))
refman = load_tsv(os.path.join(REFDIR, 'reference_construction.tsv'))

KEY = {'attJ_junction': 'attJ (chr08, excised state)', 'unit1_to_unit2_junction': 'plasmid boundary A (u1->u2)',
       'unit2_to_unit1_circular_junction': 'plasmid boundary B (origin)',
       'retained_window': 'retained state (chr06)', 'neg_attJ_shuffled': 'negative: shuffled attJ',
       'neg_ctrl_chr08_1': 'negative: GC-matched ctrl 1', 'neg_ctrl_chr08_2': 'negative: GC-matched ctrl 2',
       'neg_ctrl_chr08_3': 'negative: GC-matched ctrl 3'}

# Panel A: junction support summary（panel 参考 + full 参考）
panel_rows = []
for r in summ:
    panel_rows.append([KEY.get(r['ref'], r['ref']), r['type'] if 'type' in r else r.get('jtype', ''), r['jpos'],
                       r['SAM_loose_support_reads'], r['SAM_strict_support_reads'],
                       r['PAF_loose_mapq20'], r['PAF_strict_mapq20']])
state_rows = []
for r in sjc:
    state_rows.append([r['dataset'], r['junction'], r['loose_reads'], r['strict_reads']])
# Panel B: supporting reads（strict）
verdict_by_read = defaultdict(list)
for r in cls3:
    verdict_by_read[r['read']].append(r)
sup_rows = []
for r in cand:
    if r['pass_strict'] == 'True':
        v = (verdict_by_read.get(r['read']) or [{}])[0]
        sup_rows.append([KEY.get(r['ref'], r['ref']), r['read'], r['read_len'], r['mapq'], r['ident'],
                         r['a_left'], r['a_right'], r['AS'], r['NM'], r['clip_l'], r['clip_r'],
                         v.get('best_model', ''), v.get('delta', ''), v.get('verdict', '')])
sup_rows.sort(key=lambda x: (x[0], -int(x[7] or 0)))

with open(os.path.join(T, 'Supplementary_Table_S3_junction_read_support.tsv'), 'w', encoding='utf-8') as f:
    f.write('# Panel A1: junction support summary (10-kb junction references; SAM and PAF paths)\n')
    f.write('junction\ttype\tjunction_pos(0-based)\tloose_reads\tstrict_reads\tPAF_loose_reads\tPAF_strict_reads\n')
    for r_ in panel_rows:
        f.write('\t'.join(str(x) for x in r_) + '\n')
    f.write('\n# Panel A2: chromosome-state junction counts (full-db mapping; smbu06 = pure-retained control)\n')
    f.write('dataset\tjunction\tloose_reads(mapQ>=20,1kb)\tstrict_reads(3kb)\n')
    for r_ in state_rows:
        f.write('\t'.join(str(x) for x in r_) + '\n')
    f.write('\n# Panel B: strict supporting reads\n')
    f.write('junction\tread\tread_len\tmapQ\tidentity\tanchor_left\tanchor_right\tAS\tNM\tclip_left\tclip_right\tbest_model\tdelta_AS\tverdict\n')
    for r_ in sup_rows:
        f.write('\t'.join(str(x) for x in r_) + '\n')

try:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'A1_junction_summary'
    ws.append(['junction', 'type', 'junction_pos', 'loose_reads', 'strict_reads', 'PAF_loose', 'PAF_strict'])
    for r_ in panel_rows:
        ws.append(r_)
    ws2 = wb.create_sheet('A2_state_counts')
    ws2.append(['dataset', 'junction', 'loose_reads', 'strict_reads'])
    for r_ in state_rows:
        ws2.append(r_)
    ws3 = wb.create_sheet('B_supporting_reads')
    ws3.append(['junction', 'read', 'read_len', 'mapQ', 'identity', 'anchor_left', 'anchor_right', 'AS', 'NM',
                'clip_left', 'clip_right', 'best_model', 'delta_AS', 'verdict'])
    for r_ in sup_rows:
        ws3.append(r_)
    ws4 = wb.create_sheet('C_model_comparison')
    ws4.append(['read', 'jtype', 'read_len', 'best_model', 'best_AS', 'second_model', 'second_AS', 'delta',
                'best_ident', 'second_ident', 'cover', 'verdict'])
    for r in cls3:
        ws4.append([r[k] for k in ['read', 'jtype', 'read_len', 'best_model', 'best_AS', 'second_model',
                                   'second_AS', 'delta', 'best_ident', 'second_ident', 'cover', 'verdict']])
    ws5 = wb.create_sheet('D_depth_ratio')
    ws5.append(['region', 'target', 'ratio_primary', 'ratio_allaln', 'CI95_low', 'CI95_high'])
    for r in dbs:
        ws5.append([r['region'], r['target'], r['ratio_primary_algA'], r['ratio_allaln_algB'],
                    r['bootstrap_CI95_low'], r['bootstrap_CI95_high']])
    ws6 = wb.create_sheet('E_copy_number')
    for r in cn:
        ws6.append([r['item'], r['value'], r['detail']])
    ws7 = wb.create_sheet('F_1bp_withinread')
    ws7.append(['group', 'n_reads', 'mean_delta', 'CI95_low', 'CI95_high'])
    for r in (r1 or r3):
        ws7.append([r.get('group', ''), r.get('n_reads', r.get('n_reads_with_both_windows', '')),
                    r.get('mean_delta', ''), r.get('bootstrap_CI95_low', ''), r.get('bootstrap_CI95_high', '')])
    for r in r1p:
        ws7.append(['pair:' + r['group'], r['read'], r['run_copy1'], r['run_copy2'], r['delta(1-2)']])
    for r in r3p:
        ws7.append(['pair3:' + r['group'], r['read'], r['run_copy1'], r['run_copy2'], r['delta(1-2)']])
    ws8 = wb.create_sheet('G_multimapping')
    for r in mm:
        ws8.append([r['item'], r['value']])
    ws9 = wb.create_sheet('H_depth_regions')
    ws9.append(['region', 'target', 'n_positions', 'depth_primary_mean', 'depth_primary_median', 'depth_primary_IQR',
                'depth_allaln_mean', 'depth_allaln_median', 'depth_allaln_IQR'])
    for r in drs:
        ws9.append([r['region'], r['target'], r['n_positions'], r['depth_primary_mean'], r['depth_primary_median'],
                    r['depth_primary_IQR'], r['depth_allaln_mean'], r['depth_allaln_median'], r['depth_allaln_IQR']])
    wb.save(os.path.join(T, 'Supplementary_Table_S3_junction_read_support.xlsx'))
    print('xlsx written')
except ImportError:
    print('openpyxl unavailable')

with open(os.path.join(T, 'junction_reference_manifest.tsv'), 'w', encoding='utf-8') as f:
    f.write('reference\tlength_bp\tfasta_sha256\tconstruction\n')
    for r in refman:
        f.write('\t'.join([r['reference'], r['length_bp'], r['fasta_sha256'], r['construction']]) + '\n')

with open(os.path.join(T, 'mapping_model_comparison.tsv'), 'w', encoding='utf-8') as f:
    f.write('read\tjtype\tread_len\tbest_model\tbest_AS\tsecond_model\tsecond_AS\tdelta_AS\tbest_identity\tsecond_identity\tread_coverage\tverdict\n')
    for r in cls3:
        f.write('\t'.join([r[k] for k in ['read', 'jtype', 'read_len', 'best_model', 'best_AS', 'second_model',
                                          'second_AS', 'delta', 'best_ident', 'second_ident', 'cover', 'verdict']]) + '\n')

with open(os.path.join(T, 'ambiguous_and_negative_reads.tsv'), 'w', encoding='utf-8') as f:
    f.write('# Part 1: ambiguous candidate reads (pre-registered: delta_AS < 10 vs best alternative model)\n')
    f.write('read\tjtypes\tread_len\tbest_model\tbest_AS\tsecond_model\tsecond_AS\tdelta_AS\tbest_identity\tverdict\n')
    for r in cls3:
        if r['verdict'] == 'ambiguous':
            f.write('\t'.join([r['read'], r['jtype'], r['read_len'], r['best_model'], r['best_AS'], r['second_model'],
                               r['second_AS'], r['delta'], r['best_ident'], r['verdict']]) + '\n')
    f.write('\n# Part 2: negative-control junction candidates (all alignments on negative references)\n')
    f.write('ref\tread\tmapq\tprimary\tcovered\ta_left\ta_right\tpass_loose\tpass_strict\n')
    for r in cand:
        if r['ref'].startswith('neg_'):
            f.write('\t'.join([r['ref'], r['read'], r['mapq'], r['primary'], r['covered'],
                               str(r['a_left']), str(r['a_right']), r['pass_loose'], r['pass_strict']]) + '\n')

with open(os.path.join(T, 'chromosome_state_and_copy_number.tsv'), 'w', encoding='utf-8') as f:
    f.write('# Chromosome state counts and copy-number estimates\n')
    f.write('## junction counts (loose reads; see Supplementary_Table_S3 Panel A2 for detail)\n')
    for r in sjc:
        f.write('state_count\t%s\t%s\tloose=%s\tstrict=%s\n' % (r['dataset'], r['junction'], r['loose_reads'], r['strict_reads']))
    f.write('## copy number (tandem-repeat-corrected)\n')
    for r in cn:
        f.write('copy_number\t%s\t%s\t%s\n' % (r['item'], r['value'], r['detail']))
print('build_tables2 done')
